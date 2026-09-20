"""MANUSCRIPT.md -> MANUSCRIPT.docx.

Biçim kuralı (kullanıcı): collapsible/Word-Heading stili YOK; bölüm etiketleri düz Normal
paragrafta bold + bir punto büyük. Gövde 12pt, bölüm etiketi 13pt, başlık 14pt. Hepsi Normal
tabanlı → Word navigasyon panelinde katlanmaz. Kelime-arası em-dash zaten kaynakta yok.
"""
import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "MANUSCRIPT.md"
OUT = "MANUSCRIPT.docx"

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(12)

_tok = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")

def add_runs(par, text, base_bold=False, base_size=12):
    """Inline **bold** ve *italic* işaretlerini run'lara çevir."""
    for part in _tok.split(text):
        if not part:
            continue
        bold, italic, inner = base_bold, False, part
        if part.startswith("**") and part.endswith("**"):
            bold, inner = True, part[2:-2]
        elif part.startswith("*") and part.endswith("*"):
            italic, inner = True, part[1:-1]
        r = par.add_run(inner)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(base_size)

lines = open(SRC, encoding="utf-8").read().split("\n")
title_done = False
for raw in lines:
    line = raw.rstrip()
    if line == "" or line == "---":
        continue
    # tamamen **...** olan satır = başlık/bölüm etiketi
    if re.fullmatch(r"\*\*.+\*\*", line):
        inner = line[2:-2]
        p = doc.add_paragraph()
        if not title_done:
            add_runs(p, inner, base_bold=True, base_size=14)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_done = True
        else:
            add_runs(p, inner, base_bold=True, base_size=13)
            p.paragraph_format.space_before = Pt(10)
        continue
    if line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, line[2:])
        continue
    p = doc.add_paragraph()
    add_runs(p, line)

doc.save(OUT)
print("yazıldı:", OUT, "| paragraf:", len(doc.paragraphs))
